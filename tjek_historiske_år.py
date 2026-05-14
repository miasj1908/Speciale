"""
Tjekker hvor mange komplette historiske år hver cashflow model indeholder.
Forsøger multiple indlæsningsmetoder inkl. win32com (Excel direkte).

Krav:
    pip install xlrd openpyxl pywin32
"""

import pandas as pd
import xlrd
import openpyxl
import os
import tempfile
from datetime import datetime

# ── INDSTILLINGER ─────────────────────────────────────────────────────────────

MATCHEDE_PAR_STI = r"C:\Users\miasj\PythonProjects\matchede_par_final.csv"

LOKAL_ROD_STI = r"C:\Users\miasj\CBS - Copenhagen Business School\Caroline Thøisen Larsen - Data - Forecasting downloadet"

SHAREPOINT_PRAEFIKS = "https://studentcbs-my.sharepoint.com/personal/ctl_acc_cbs_dk/Documents/Data - Forecasting downloadet"

# ─────────────────────────────────────────────────────────────────────────────


def konverter_sti(sharepoint_sti):
    if pd.isna(sharepoint_sti):
        return None
    sti = str(sharepoint_sti).strip()
    if sti.startswith("http"):
        sti = sti.replace(SHAREPOINT_PRAEFIKS, LOKAL_ROD_STI)
        sti = sti.replace("/", "\\")
        sti = sti.replace("Rådata", "Radata")
    return sti


def find_historiske_aar_fra_rows(raekker):
    """Finder historiske år fra en liste af rækker."""
    offset_raekke = None
    aar_raekke = None

    for row in raekker:
        row = [v for v in row if v is not None]
        if -10 in row or -10.0 in row:
            offset_raekke = row
        if any(isinstance(v, (int, float)) and 2000 < v < 2035 for v in row):
            aar_raekke = row

    if offset_raekke is None or aar_raekke is None:
        return 0, False, []

    historiske_aar = []
    for offset, aar in zip(offset_raekke, aar_raekke):
        if isinstance(offset, (int, float)) and isinstance(aar, (int, float)):
            if -5 <= int(offset) <= -1:
                historiske_aar.append(int(aar))

    if not historiske_aar:
        return 0, False, []

    historiske_aar_sorted = sorted(historiske_aar)
    er_sammenhaengende = all(
        historiske_aar_sorted[i+1] - historiske_aar_sorted[i] == 1
        for i in range(len(historiske_aar_sorted) - 1)
    )

    return len(historiske_aar), er_sammenhaengende, historiske_aar_sorted


def find_inputs_navn(sheet_navne):
    """Find Inputs sheet navn."""
    for navn in ['Inputs', 'inputs', 'INPUTS']:
        if navn in sheet_navne:
            return navn
    return None


def indlaes_xlrd(filsti):
    """Standard xlrd indlaesning."""
    try:
        wb = xlrd.open_workbook(filsti)
        inputs_navn = find_inputs_navn(wb.sheet_names())
        if not inputs_navn:
            return None
        sheet = wb.sheet_by_name(inputs_navn)
        raekker = [sheet.row_values(i) for i in range(min(10, sheet.nrows))]
        return raekker
    except Exception:
        return None


def indlaes_xlrd_encoding(filsti, encoding):
    """xlrd med specifik encoding."""
    try:
        wb = xlrd.open_workbook(filsti, encoding_override=encoding)
        inputs_navn = find_inputs_navn(wb.sheet_names())
        if not inputs_navn:
            return None
        sheet = wb.sheet_by_name(inputs_navn)
        raekker = [sheet.row_values(i) for i in range(min(10, sheet.nrows))]
        return raekker
    except Exception:
        return None


def indlaes_xlrd_ignore(filsti):
    """xlrd med ignore_workbook_corruption."""
    try:
        wb = xlrd.open_workbook(filsti, ignore_workbook_corruption=True)
        inputs_navn = find_inputs_navn(wb.sheet_names())
        if not inputs_navn:
            return None
        sheet = wb.sheet_by_name(inputs_navn)
        raekker = [sheet.row_values(i) for i in range(min(10, sheet.nrows))]
        return raekker
    except Exception:
        return None


def indlaes_win32com(filsti):
    """
    Aabner filen med Excel via win32com.
    Konverterer til xlsx i temp-mappe og laesser med openpyxl.
    """
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        temp_dir = tempfile.mkdtemp()
        temp_xlsx = os.path.join(temp_dir, 'temp_converted.xlsx')

        excel = win32com.client.Dispatch('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False

        wb_excel = excel.Workbooks.Open(os.path.abspath(filsti))
        wb_excel.SaveAs(temp_xlsx, FileFormat=51)  # 51 = xlsx
        wb_excel.Close(False)
        excel.Quit()

        pythoncom.CoUninitialize()

        # Laes den konverterede fil med openpyxl
        wb = openpyxl.load_workbook(temp_xlsx, read_only=True, data_only=True)
        inputs_navn = find_inputs_navn(wb.sheetnames)
        if not inputs_navn:
            return None
        sheet = wb[inputs_navn]
        raekker = [list(row) for row in sheet.iter_rows(
            min_row=1, max_row=10, values_only=True)]

        # Ryd op
        try:
            os.remove(temp_xlsx)
            os.rmdir(temp_dir)
        except Exception:
            pass

        return raekker

    except Exception:
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
        return None


def indlaes_fil(filsti, brug_win32=False):
    """
    Forsøger multiple metoder.
    brug_win32=True aktiverer win32com som sidste udvej.
    """
    # Metode 1: Standard xlrd
    raekker = indlaes_xlrd(filsti)
    if raekker is not None:
        return raekker, 'xlrd_standard'

    # Metode 2: UTF-16
    raekker = indlaes_xlrd_encoding(filsti, 'utf-16')
    if raekker is not None:
        return raekker, 'xlrd_utf16'

    # Metode 3: Latin-1
    raekker = indlaes_xlrd_encoding(filsti, 'latin-1')
    if raekker is not None:
        return raekker, 'xlrd_latin1'

    # Metode 4: Ignore corruption
    raekker = indlaes_xlrd_ignore(filsti)
    if raekker is not None:
        return raekker, 'xlrd_ignore'

    # Metode 5: win32com (Excel direkte)
    if brug_win32:
        raekker = indlaes_win32com(filsti)
        if raekker is not None:
            return raekker, 'win32com_excel'

    return None, 'alle_metoder_fejlede'


def koer_tjek():
    """Hovedfunktion der scanner alle matchede filer."""

    print("Indlaeser matchede_par_final.csv...")
    df = pd.read_csv(MATCHEDE_PAR_STI, sep=';', encoding='utf-8-sig',
                     on_bad_lines='skip')
    print(f"Fandt {len(df)} matchede par")
    print("Bruger win32com (Excel) som backup for fejlende filer\n")

    resultater = []
    fejl = []
    metode_taeller = {}
    win32_taeller = 0

    for idx, raekke in df.iterrows():
        cf_mappe = konverter_sti(raekke['cf_mappe'])
        cf_fil = raekke['cf_fil']

        if cf_mappe is None:
            fejl.append({'fil': cf_fil, 'fejl': 'Ingen mappe-sti'})
            continue

        filsti = os.path.join(cf_mappe, cf_fil)

        if idx % 100 == 0:
            print(f"Behandler fil {idx}/{len(df)}: {cf_fil}")

        if not os.path.exists(filsti):
            fejl.append({'fil': cf_fil, 'fejl': 'Fil ikke fundet'})
            continue

        # Prøv først uden win32com
        raekker, metode = indlaes_fil(filsti, brug_win32=False)

        # Hvis det fejler, prøv win32com
        if raekker is None:
            win32_taeller += 1
            if win32_taeller % 10 == 0:
                print(f"  Bruger Excel til fil {win32_taeller}: {cf_fil}")
            raekker, metode = indlaes_fil(filsti, brug_win32=True)

        if raekker is None:
            fejl.append({'fil': cf_fil, 'fejl': metode})
            continue

        metode_taeller[metode] = metode_taeller.get(metode, 0) + 1

        antal_aar, sammenhaengende, aar_liste = find_historiske_aar_fra_rows(raekker)

        resultater.append({
            'selskab':         raekke['selskab'],
            'sektor':          raekke['sektor'],
            'cf_fil':          cf_fil,
            'rapport_dato':    raekke['rapport_dato'],
            'antal_hist_aar':  antal_aar,
            'sammenhaengende': sammenhaengende,
            'aar_liste':       str(aar_liste),
            'metode':          metode,
        })

    df_res = pd.DataFrame(resultater)

    print("\n" + "="*50)
    print("RESULTATER")
    print("="*50)
    print(f"\nSucces:  {len(resultater)} filer analyseret")
    print(f"Fejl:    {len(fejl)} filer kunne ikke laeses")

    print("\n--- Indlaesningsmetoder brugt ---")
    for metode, antal in sorted(metode_taeller.items()):
        print(f"  {metode}: {antal}")

    if len(df_res) > 0:
        print("\n--- Fordeling af historiske aar ---")
        print(df_res['antal_hist_aar'].value_counts().sort_index().to_string())

        print("\n--- Kvalificerede filer per krav ---")
        for krav in [5, 4, 3]:
            antal = len(df_res[
                (df_res['antal_hist_aar'] >= krav) &
                (df_res['sammenhaengende'] == True)
            ])
            print(f"  Mindst {krav} sammenhaengende aar: {antal} filer ({antal/len(df_res)*100:.1f}%)")

        print("\n--- Fordeling per sektor (5 sammenhaengende aar) ---")
        kvalificerede = df_res[
            (df_res['antal_hist_aar'] >= 5) &
            (df_res['sammenhaengende'] == True)
        ]
        print(kvalificerede['sektor'].value_counts().to_string())

    # Gem resultater
    try:
        output_sti = r"C:\Users\miasj\PythonProjects\historiske_aar_tjek.csv"
        df_res.to_csv(output_sti, index=False)
        print(f"\nDetaljeret resultat gemt til: {output_sti}")
    except Exception as e:
        print(f"\nKunne ikke gemme resultat: {e}")

    try:
        fejl_sti = r"C:\Users\miasj\PythonProjects\fejl_log_ny.csv"
        pd.DataFrame(fejl).to_csv(fejl_sti, index=False)
        print(f"Fejlliste gemt til: {fejl_sti}")
    except Exception as e:
        print(f"Kunne ikke gemme fejlliste: {e}")

    return df_res


if __name__ == "__main__":
    start = datetime.now()
    df_res = koer_tjek()
    slut = datetime.now()
    minutter = (slut - start).seconds // 60
    sekunder = (slut - start).seconds % 60
    print(f"\nTid brugt: {minutter} minutter og {sekunder} sekunder")
